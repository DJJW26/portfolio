"""Resume sync script: updates Divij Jhunjhunwala.docx (CloudSim entries, idempotent).

NOTE: PDF export is owned by Divij now — export from Word yourself as
Divij-Jhunjhunwala-Resume.pdf and copy it to client/public/. This script
no longer writes that file (its PDF output goes to a -generated name).

Usage:  python build_resume.py
Outputs: Divij Jhunjhunwala.docx (updated), Divij-Jhunjhunwala-Resume-generated.pdf
"""
import docx
from fpdf import FPDF

DOCX_PATH = "Divij Jhunjhunwala.docx"
PDF_PATH = "Divij-Jhunjhunwala-Resume-generated.pdf"

CLOUDSIM_TITLE = "Energy-Aware VM Placement (CloudSim)"
CLOUDSIM_BULLETS = [
    "Implemented PABFD and Round-Robin VM allocation policies in CloudSim 3.0.3 to reproduce Beloglazov et al. (2012)",
    "PABFD cut datacenter energy use by 48.6% with zero SLA violation; MBFD with migration cut 57.7% at ~3% SLA violation",
    "Code: github.com/DJJW26/CloudSim",
]

LINKS = {
    "email": "mailto:divijjjw@gmail.com",
    "linkedin": "https://www.linkedin.com/in/divij-jhunjhunwala-418b3027a/",
    "github": "https://github.com/DJJW26",
}


def update_docx():
    doc = docx.Document(DOCX_PATH)
    texts = [p.text for p in doc.paragraphs]

    # 1. Add CloudSim to Dev Tools line (idempotent)
    for p in doc.paragraphs:
        if p.text.startswith("Dev Tools:") and "CloudSim" not in p.text:
            p.text = p.text.rstrip() + ", CloudSim"
            break

    # 2. Append CloudSim project at the end (idempotent)
    if not any(CLOUDSIM_TITLE in t for t in texts):
        # match the plain style of surrounding paragraphs
        style = doc.paragraphs[-1].style if doc.paragraphs else "Normal"
        doc.add_paragraph("", style=style)
        doc.add_paragraph(CLOUDSIM_TITLE, style=style)
        for b in CLOUDSIM_BULLETS:
            doc.add_paragraph(b, style=style)

    doc.save(DOCX_PATH)
    return [p.text for p in doc.paragraphs]


def clean(s):
    """Map unicode punctuation to latin-1 safe equivalents (core fonts)."""
    return (
        s.replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2022", "-")
        .replace("\u00a0", " ")
    )


class ResumePDF(FPDF):
    pass


def build_pdf(paragraphs):
    paragraphs = [clean(p) for p in paragraphs]
    pdf = ResumePDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    W = pdf.w - 20  # 10mm margins each side
    pdf.set_left_margin(10)
    pdf.set_right_margin(10)

    for i, text in enumerate(paragraphs):
        t = text.strip()
        if not t:
            pdf.ln(2)
            continue
        if i == 0:  # Name
            pdf.set_font("Helvetica", "B", 22)
            pdf.cell(W, 10, t, new_x="LMARGIN", new_y="NEXT", align="C")
        elif i == 1:  # contact line with links
            pdf.set_font("Helvetica", "", 10)
            parts = [p.strip() for p in t.split("|")]
            # render segments, attaching links where known
            pdf.set_x(10)
            for j, part in enumerate(parts):
                if j:
                    pdf.write(5, "  |  ")
                low = part.lower()
                if "@" in part:
                    pdf.set_text_color(0, 0, 180)
                    pdf.write(5, part, link=LINKS["email"])
                    pdf.set_text_color(0, 0, 0)
                elif low == "linkedin":
                    pdf.set_text_color(0, 0, 180)
                    pdf.write(5, part, link=LINKS["linkedin"])
                    pdf.set_text_color(0, 0, 0)
                elif low == "github":
                    pdf.set_text_color(0, 0, 180)
                    pdf.write(5, part, link=LINKS["github"])
                    pdf.set_text_color(0, 0, 0)
                else:
                    pdf.write(5, part)
            pdf.ln(7)
        elif t in ("Experience:", "Projects:"):
            pdf.set_font("Helvetica", "B", 13)
            pdf.ln(1)
            pdf.cell(W, 7, t, new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(0, 180, 200)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(2)
        elif any(t.startswith(k) for k in ("Languages:", "Web Technologies:", "Databases:", "Dev Tools:", "Core Concepts:")):
            head, _, rest = t.partition(":")
            pdf.set_font("Helvetica", "B", 10)
            pdf.write(5, head + ":")
            pdf.set_font("Helvetica", "", 10)
            pdf.write(5, rest)
            pdf.ln(6)
        elif t in ("Speech Evaluation Web Application", "Parking Lot application", CLOUDSIM_TITLE,
                   "Vice-President, Computer Club - St. Xavier's Collegiate School",
                   "Vice-President, Computer Club - St. Xavier\u2019s Collegiate School"):
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(W, 6, t, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(W, 5, t)

    pdf.output(PDF_PATH)
    print(f"Wrote {PDF_PATH}")


if __name__ == "__main__":
    paras = update_docx()
    print("DOCX updated.")
    build_pdf(paras)
